"""
Step 1: Read the Eurostat sts_trtu_m Excel file (48 sheets) and make ONE clean long table.
Usage:  python 01_build_clean_data.py  path/to/your_file.xlsx
Output: sts_retail_clean_long.csv  (columns: country, indicator, nace, s_adj, unit, month, value, flag)
"""
import sys, warnings
warnings.filterwarnings("ignore")
import pandas as pd
from openpyxl import load_workbook

SRC = sys.argv[1] if len(sys.argv) > 1 else "sts_trtu_m__custom_22985693_spreadsheet.xlsx"
IND  = {"Volume of sales": "VOL_SLS", "Net turnover": "NETTUR"}
NACE = {"Retail trade, except of motor vehicles and motorcycles": "G47",
        "Retail sale of food, beverages and tobacco": "G47_FOOD",
        "Retail sale of non-food products (including fuel)": "G47_NFOOD",
        "Retail sale via mail order houses or via Internet": "G4791"}
ADJ  = {"Unadjusted data (i.e. neither seasonally adjusted nor calendar adjusted data)": "NSA",
        "Calendar adjusted data, not seasonally adjusted data": "CA",
        "Seasonally and calendar adjusted data": "SCA"}
UNIT = {"Index, 2021=100": "I21",
        "Percentage change compared to same period in previous year": "PCH_SM"}

wb = load_workbook(SRC, data_only=True)       # NOT read_only: Eurostat files report wrong sizes in read-only mode
out = []
for name in wb.sheetnames:
    if not name.startswith("Sheet "):
        continue
    rows = list(wb[name].iter_rows(values_only=True))
    ind, nace = IND[rows[5][2]], NACE[rows[6][2]]
    adj, unit = ADJ[rows[7][2]], UNIT[rows[8][2]]
    header = rows[10]                                     # TIME row: month, (flag column), month, ...
    months = [(i, header[i]) for i in range(1, len(header)) if header[i]]
    for r in rows[12:]:                                   # country rows start at row 13
        if not r[0]:
            break
        for i, m in months:
            out.append((r[0], ind, nace, adj, unit, m, r[i], r[i + 1] if i + 1 < len(r) else None))

df = pd.DataFrame(out, columns=["country", "indicator", "nace", "s_adj", "unit", "month", "value", "flag"])
df["value"] = pd.to_numeric(df["value"], errors="coerce")   # ':' (missing) becomes NaN
df["month"] = pd.to_datetime(df["month"])
df["flag"]  = df["flag"].fillna("")
df.loc[(df.unit == "I21") & (df.value <= 0), "value"] = float("nan")   # Eurostat placeholder zeros (e.g. Denmark online 2010-14)
df = df.dropna(subset=["value"])
df.to_csv("sts_retail_clean_long.csv", index=False)
print("Saved sts_retail_clean_long.csv:", df.shape)