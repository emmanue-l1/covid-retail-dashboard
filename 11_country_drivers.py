"""
Step 11: Why did countries differ?  (association study, not proof of cause)
Run : python 11_country_drivers.py        (needs steps 06 and 08, and the downloaded files, see below)
Files (put them in the project folder or in a folder called extra_data):
      lockdown*.csv       Our World in Data, "Estimated daily excess deaths per 100,000 people" (central estimate) and
                          confirmed COVID-19 deaths per 100,000. (The file name says lockdown, but the content is deaths.)
      tec00114*.xlsx      Eurostat, GDP per capita in PPS (EU27 = 100)
      isoc_ec_ibuy*.xlsx  Eurostat, internet purchases by individuals (sheet "in the 12 months", % of individuals)
      any csv with a stringency column, OPTIONAL: the Oxford COVID-19 Government Response Tracker file (column StringencyIndex_Average, 0 to 100)
                          or the Our World in Data stringency chart. It is the only measure here of how strict the rules were.
                          Oxford does not cover Montenegro and North Macedonia, so those two drop out of the regressions.
Out : results/drivers/country_features.csv, correlations.csv, regression.csv, regression_stability.csv, predictor_correlations.csv, wave_correlations.csv

What is compared: for every country, the average COVID gap in retail sales (from step 06) against country facts:
  1. excess deaths     = average estimated excess deaths per 100,000 people per day, Mar 2020 - Dec 2021 (health burden; better than confirmed
                         deaths, because countries counted deaths differently)
  2. confirmed deaths  = average confirmed COVID-19 deaths per 100,000 per day, same period (shown in the correlations only)
  3. online shopping   = % of people who bought online in the 12 months before the survey, 2019 (before COVID)
  4. income            = GDP per capita in PPS, 2019 (EU27 = 100)
  5. strictness        = average stringency index, Mar 2020 - Dec 2021 (only if the optional file is there)
Methods and why:
  - Spearman rank correlation: simple, no assumption about straight lines, robust to outliers (e.g. Greece online).
  - Multiple regression on standardised predictors (one unit = one standard deviation) with HC3 robust standard errors.
    Predictors: excess deaths, online shopping, income and (if available) strictness. It shows what each fact adds when the others
    are held fixed. With only 20-30 countries it is a guide, not proof.
  - Drop-one-country check: does the sign of each coefficient survive removing any single country?
"""
import os, glob, warnings
warnings.filterwarnings("ignore")
import numpy as np, pandas as pd
from scipy import stats
import openpyxl

OUT = "results/drivers/"; os.makedirs(OUT, exist_ok=True)
NAMES = {"G47": "Total retail", "G47_FOOD": "Food, drink, tobacco", "G47_NFOOD": "Non-food (incl. fuel)", "G4791": "Online / mail order"}
ISO = {"Albania": "ALB", "Austria": "AUT", "Belgium": "BEL", "Bosnia and Herzegovina": "BIH", "Bulgaria": "BGR", "Croatia": "HRV",
       "Cyprus": "CYP", "Czechia": "CZE", "Denmark": "DNK", "Estonia": "EST", "Finland": "FIN", "France": "FRA", "Germany": "DEU",
       "Greece": "GRC", "Hungary": "HUN", "Ireland": "IRL", "Italy": "ITA", "Latvia": "LVA", "Lithuania": "LTU", "Luxembourg": "LUX",
       "Malta": "MLT", "Moldova": "MDA", "Montenegro": "MNE", "Netherlands": "NLD", "North Macedonia": "MKD", "Norway": "NOR",
       "Poland": "POL", "Portugal": "PRT", "Romania": "ROU", "Serbia": "SRB", "Slovakia": "SVK", "Slovenia": "SVN", "Spain": "ESP",
       "Sweden": "SWE", "Switzerland": "CHE", "Türkiye": "TUR", "Turkey": "TUR", "United Kingdom": "GBR", "Iceland": "ISL", "Czech Republic": "CZE"}
WAVES = [("First lockdown", "2020-03-01", "2020-06-30"), ("Summer 2020", "2020-07-01", "2020-10-31"),
         ("Winter wave", "2020-11-01", "2021-04-30"), ("Reopening 2021", "2021-05-01", "2021-12-31")]


def find(pattern):
    for folder in [".", "extra_data"]:
        hit = sorted(glob.glob(os.path.join(folder, pattern)))
        if hit:
            return hit[0]
    raise SystemExit(f"Could not find a file like {pattern}. Put the three downloaded files in the project folder or in a folder called extra_data.")


def eurostat_year(path, sheet, year):
    """Read a Eurostat Excel sheet (years in one row, a flag column after every value) and return {country: value} for the year."""
    ws = openpyxl.load_workbook(path, data_only=True)[sheet]
    rows = list(ws.iter_rows(values_only=True))
    h = next(i for i, r in enumerate(rows) if r and r[0] == "TIME")
    col = list(rows[h]).index(str(year))
    out = {}
    for r in rows[h + 2:]:
        if r[0] is None or r[0] in ("Special value", "Observation flags:"):
            break
        try:
            out[str(r[0]).replace("*", "").strip()] = float(r[col])
        except (TypeError, ValueError):
            pass
    return out


# ------------------------------------------------------------------ 1. covariates per country
lock = pd.read_csv(find("lockdown*.csv"), parse_dates=["Day"]).rename(columns={"Central estimate": "excess", "Confirmed COVID-19 deaths (per 100,000)": "deaths"})
lock = lock[lock.Code.notna() & (lock.Code.str.len() == 3)]
cov = {}
w = lock[(lock.Day >= "2020-03-01") & (lock.Day <= "2021-12-31")]
cov["excess_deaths_per_100k_per_day"] = w.groupby("Code").excess.mean()
cov["confirmed_deaths_per_100k_per_day"] = w.groupby("Code").deaths.mean()
for nm, a, b in WAVES:
    ww = lock[(lock.Day >= a) & (lock.Day <= b)]
    cov["excess_" + nm] = ww.groupby("Code").excess.mean()
def find_stringency():
    """Look for a csv with a stringency column: Oxford OxCGRT (StringencyIndex_Average) or the Our World in Data version."""
    for folder in [".", "extra_data"]:
        for f in sorted(glob.glob(os.path.join(folder, "*.csv"))):
            try:
                head = pd.read_csv(f, nrows=2).columns
            except Exception:
                continue
            if "StringencyIndex_Average" in head or "stringency_index" in head:
                return f
    return None


SFILE = find_stringency()
HAS_STR = SFILE is not None
if HAS_STR:
    sg = pd.read_csv(SFILE, low_memory=False)
    if "StringencyIndex_Average" in sg.columns:                      # Oxford OxCGRT full file
        sg = sg[~sg.Date.astype(str).str.startswith("#")]            # the file has a tag row under the header
        sg["Day"] = pd.to_datetime(sg.Date.astype(str), format="%Y%m%d"); sg["Code"] = sg.CountryCode
        sg["stringency"] = pd.to_numeric(sg.StringencyIndex_Average, errors="coerce")
        print("Stringency file:", SFILE, "(Oxford COVID-19 Government Response Tracker, StringencyIndex_Average)")
    else:                                                             # Our World in Data stringency chart
        sg["Day"] = pd.to_datetime(sg.Day); sg["stringency"] = sg.stringency_index
        print("Stringency file:", SFILE, "(Our World in Data)")
    sg = sg[sg.Code.notna() & (sg.Code.astype(str).str.len() == 3)]
    cov["stringency_index"] = sg[(sg.Day >= "2020-03-01") & (sg.Day <= "2021-12-31")].groupby("Code").stringency.mean()
    for nm, a_, b_ in WAVES:
        cov["stringency_" + nm] = sg[(sg.Day >= a_) & (sg.Day <= b_)].groupby("Code").stringency.mean()
else:
    print("(No stringency file found: the study runs without a measure of how strict the rules were.)")
sheets = openpyxl.load_workbook(find("isoc_ec_ibuy*.xlsx"), data_only=True)
sh = next(s.title for s in sheets.worksheets if any("in the 12 months" in str(c) for r in s.iter_rows(min_row=1, max_row=12, values_only=True) for c in r if c))
ib = eurostat_year(find("isoc_ec_ibuy*.xlsx"), sh, 2019)
gdp = eurostat_year(find("tec00114*.xlsx"), "Sheet 1", 2019)
cov["online_shopping_2019_pct"] = pd.Series({ISO[k]: v for k, v in ib.items() if k in ISO})
cov["gdp_per_capita_pps_2019"] = pd.Series({ISO[k]: v for k, v in gdp.items() if k in ISO})
C = pd.DataFrame(cov); C.index.name = "iso3"
print(f"Covariates found: deaths {C.excess_deaths_per_100k_per_day.notna().sum()} countries, online shopping {C.online_shopping_2019_pct.notna().sum()}, "
      f"GDP {C.gdp_per_capita_pps_2019.notna().sum()}")

# ------------------------------------------------------------------ 2. outcomes: average gap per country (same basis as the story page)
imp = pd.read_csv("results/final/impact_by_series.csv")
has_car = ((imp.adj == "CAR") & (imp.nace == "G47_NFOOD")).any()
PRIM = {n: ("CAR" if n == "G47_NFOOD" and has_car else "SCA") for n in NAMES}
y = imp[(imp.months_observed >= 22) & (imp.adj == imp.nace.map(PRIM))][["nace", "country", "avg_gap_pct_2020_21"]].copy()
y["iso3"] = y.country.map(ISO)
wv = pd.read_csv("results/extra/gap_by_wave_country.csv")
wv["iso3"] = wv.country.map(ISO)
D = y.merge(C, left_on="iso3", right_index=True, how="left")
D["category"] = D.nace.map(NAMES)
D.round(3).to_csv(OUT + "country_features.csv", index=False)

PRED = ["excess_deaths_per_100k_per_day", "online_shopping_2019_pct", "gdp_per_capita_pps_2019"] + (["stringency_index"] if HAS_STR else [])
ALLC = PRED + ["confirmed_deaths_per_100k_per_day"]
LAB = {"excess_deaths_per_100k_per_day": "Excess deaths", "confirmed_deaths_per_100k_per_day": "Confirmed COVID deaths",
       "online_shopping_2019_pct": "Online shopping 2019", "gdp_per_capita_pps_2019": "Income (GDP per capita)",
       "stringency_index": "Rules strictness (stringency index)"}

# ------------------------------------------------------------------ 3. correlations
rows = []
for n, g in D.groupby("nace"):
    for p in ALLC:
        x = g[["avg_gap_pct_2020_21", p]].dropna()
        if len(x) >= 8:
            r, pv = stats.spearmanr(x[p], x.avg_gap_pct_2020_21)
            rows.append((NAMES[n], LAB[p], len(x), r, pv))
cor = pd.DataFrame(rows, columns=["category", "factor", "countries", "spearman_rho", "p_value"]).round(3)
cor.to_csv(OUT + "correlations.csv", index=False)


# ------------------------------------------------------------------ 4. regression with robust (HC3) standard errors
def ols(X, yv):
    n, k = X.shape
    XtXi = np.linalg.inv(X.T @ X)
    b = XtXi @ X.T @ yv
    e = yv - X @ b
    h = np.einsum("ij,jk,ik->i", X, XtXi, X)
    meat = X.T @ np.diag(e ** 2 / (1 - h) ** 2) @ X
    V = XtXi @ meat @ XtXi
    se = np.sqrt(np.diag(V))
    pv = 2 * stats.t.sf(np.abs(b / se), n - k)
    r2 = 1 - (e @ e) / ((yv - yv.mean()) @ (yv - yv.mean()))
    return b, se, pv, r2, 1 - (1 - r2) * (n - 1) / (n - k)


reg, stab, pc_rows = [], [], []
for n, g in D.groupby("nace"):
    g = g.dropna(subset=PRED + ["avg_gap_pct_2020_21"]).reset_index(drop=True)
    if len(g) < 12:
        continue
    Z = (g[PRED] - g[PRED].mean()) / g[PRED].std()
    X = np.column_stack([np.ones(len(g)), Z.values]); yv = g.avg_gap_pct_2020_21.values
    b, se, pv, r2, ar2 = ols(X, yv)
    for i, p in enumerate(PRED):
        reg.append((NAMES[n], LAB[p], len(g), b[i + 1], se[i + 1], pv[i + 1], r2, ar2))
        signs = []
        for j in range(len(g)):                               # drop one country at a time
            keep = np.arange(len(g)) != j
            Zk = (g.loc[keep, PRED] - g.loc[keep, PRED].mean()) / g.loc[keep, PRED].std()
            bk = ols(np.column_stack([np.ones(keep.sum()), Zk.values]), yv[keep])[0]
            signs.append(bk[i + 1])
        stab.append((NAMES[n], LAB[p], len(g), b[i + 1], min(signs), max(signs), np.mean(np.sign(signs) == np.sign(b[i + 1])) * 100))
    cm = Z.corr()
    for a in PRED:
        for c in PRED:
            if a < c:
                pc_rows.append((NAMES[n], LAB[a], LAB[c], cm.loc[a, c]))
reg = pd.DataFrame(reg, columns=["category", "factor", "countries", "coef_pp_per_SD", "robust_se", "p_value", "r2", "adj_r2"]).round(3)
reg.to_csv(OUT + "regression.csv", index=False)
stab = pd.DataFrame(stab, columns=["category", "factor", "countries", "coef_pp_per_SD", "coef_min_drop_one", "coef_max_drop_one",
                                   "pct_drop_one_same_sign"]).round(2)
stab.to_csv(OUT + "regression_stability.csv", index=False)
pd.DataFrame(pc_rows, columns=["category", "factor_a", "factor_b", "correlation"]).round(2).to_csv(OUT + "predictor_correlations.csv", index=False)

# ------------------------------------------------------------------ 5. by wave: gap in the wave vs the lockdown score in the same wave
wr = []
for nm, _, _ in WAVES:
    for fac, pre in [("Excess deaths", "excess_")] + ([("Rules strictness", "stringency_")] if HAS_STR else []):
        for n, g in wv[wv.wave == nm].groupby("nace"):
            x = g.merge(C[[pre + nm]], left_on="iso3", right_index=True).dropna(subset=[pre + nm, "avg_gap_pct"])
            if len(x) >= 8:
                r, pv = stats.spearmanr(x[pre + nm], x.avg_gap_pct)
                wr.append((NAMES[n], nm, fac, len(x), r, pv))
wr = pd.DataFrame(wr, columns=["category", "wave", "factor", "countries", "spearman_rho", "p_value"]).round(3)
wr.to_csv(OUT + "wave_correlations.csv", index=False)

pd.set_option("display.width", 220); pd.set_option("display.max_columns", 20)
print("\nSPEARMAN correlation between the average gap and each factor (negative = more of the factor, bigger fall):")
print(cor.pivot(index="category", columns="factor", values="spearman_rho").to_string())
print("\np-values (below 0.05 = unlikely to be chance, but remember about 16 tests):")
print(cor.pivot(index="category", columns="factor", values="p_value").to_string())
print("\nREGRESSION: change in the gap (percentage points) for +1 standard deviation of each factor, others held fixed:")
print(reg.pivot(index="category", columns="factor", values="coef_pp_per_SD").to_string())
print("p-values:"); print(reg.pivot(index="category", columns="factor", values="p_value").to_string())
print("R2 / adj R2:"); print(reg.groupby("category")[["countries", "r2", "adj_r2"]].first().to_string())
print("\nSign stays the same when any single country is dropped (% of drops):")
print(stab.pivot(index="category", columns="factor", values="pct_drop_one_same_sign").to_string())
print("\nBy wave, Spearman (gap in the wave vs the same factor in the same wave):")
print(wr.pivot_table(index=["factor", "category"], columns="wave", values="spearman_rho").to_string())
print("\nDone. Files in", OUT)